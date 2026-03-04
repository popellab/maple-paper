"""Post-process pandoc docx: number Figure/Table captions and style them."""
import sys
from docx import Document
from docx.shared import Pt


def postprocess(path):
    doc = Document(path)
    fig_num = 0
    tab_num = 0

    for p in doc.paragraphs:
        if p.style.name == "Image Caption":
            fig_num += 1
            prefix = f"Figure {fig_num}. "
            run = p.runs[0] if p.runs else p.add_run()
            old_text = run.text
            p.clear()
            r1 = p.add_run(prefix)
            r1.bold = True
            r1.font.size = Pt(10)
            r1.font.name = "Arial"
            r2 = p.add_run(old_text)
            r2.font.size = Pt(10)
            r2.font.name = "Arial"

        elif p.style.name == "Table Caption":
            tab_num += 1
            prefix = f"Table {tab_num}. "
            run = p.runs[0] if p.runs else p.add_run()
            old_text = run.text
            p.clear()
            r1 = p.add_run(prefix)
            r1.bold = True
            r1.font.size = Pt(10)
            r1.font.name = "Arial"
            r2 = p.add_run(old_text)
            r2.font.size = Pt(10)
            r2.font.name = "Arial"

    # Bold header rows in tables
    for table in doc.tables:
        for cell in table.rows[0].cells:
            for p in cell.paragraphs:
                for run in p.runs:
                    run.bold = True

    doc.save(path)
    print(f"Post-processed {path}: {fig_num} figures, {tab_num} tables")


if __name__ == "__main__":
    for path in sys.argv[1:]:
        postprocess(path)
